#!/usr/bin/env python3
"""Build a 16:9 PPTX from one PNG per slide.

The HTML deck is the design source; this helper only places the rendered
frames. Each slide is a full-bleed picture and can optionally carry a speaker
note, so the PPTX opens identically in PowerPoint, Keynote, and LibreOffice.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


def main() -> int:
    if len(sys.argv) != 2:
        print("usage: pptx_from_images.py <request.json>", file=sys.stderr)
        return 2
    request = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    images = [Path(path) for path in request["images"]]
    output = Path(request["output"])
    if not images:
        print("no images to export", file=sys.stderr)
        return 2

    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]
    notes = request.get("notes", [])

    for index, image in enumerate(images):
        slide = presentation.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(image),
            0,
            0,
            width=presentation.slide_width,
            height=presentation.slide_height,
        )
        note = notes[index] if index < len(notes) else ""
        if note:
            slide.notes_slide.notes_text_frame.text = str(note)

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(output))
    print(json.dumps({"ok": True, "output": str(output), "slides": len(images)}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
